"""Stage 1, script 50: join everything that needs a place on a map.

Run from the project root, AFTER 10, 20 and 30:   python prep/50_lookups.py

This is where the second and third data sources meet the first:
  * World Bank population  -> students per 100,000 people     (Map 2)
  * ABS state population   -> students per 1,000 residents     (Map 4)
  * coordinates in geo/    -> circles, flow lines, campuses    (Maps 1, 3, 4)

The hard part is names. The Department of Education says "Vietnam", the
World Bank says "Viet Nam", and the map file says "Vietnam" again. Every
country gets an ISO3 code (VNM) here, once, and every later join uses the
code rather than the name.

Inputs
------
  data/nationality_by_year.csv, data/nationality_region.csv,
  data/state_by_year.csv, data/nationality_state_2025.csv,
  data/he_institution_2024.csv                          (from scripts 10-30)
  raw/worldbank_population.csv, raw/worldbank_country_metadata.csv
  raw/31010do001_202512.xlsx                            (ABS)
  geo/country_points.csv, geo/au_state_capitals.csv,
  geo/au_university_locations.csv                       (see geo/README.md)

Outputs
-------
  country_lookup.csv          one row per Dept. of Education nationality name   (reference)
  country_2025.csv            one row per country: enrolments, population, rate -> Maps 1, 2
  flows_2025.csv              top origins -> states, with both endpoints       -> Map 3
  state_2025.csv              enrolments, ABS population, rate, capital        -> Map 4 fill
  he_institution_geo_2024.csv providers with campus coordinates                -> Map 4 + beeswarm
"""
from __future__ import annotations

import sys
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent.parent
RAW, DATA, GEO = ROOT / "raw", ROOT / "data", ROOT / "geo"
YEAR = 2025
POP_YEAR = "2025"

# ------------------------------------------------------------ name fixes
# Department of Education name -> ISO3, where the name doesn't match the
# World Bank or the map file exactly. None = not a country (or no longer one)
# and deliberately left off the maps.
DOE_TO_ISO3 = {
    "Bahamas": "BHS", "Congo, Democratic Republic of": "COD", "Zaire": "COD",
    "Congo, Republic of": "COG", "Czech Republic": "CZE", "East Timor": "TLS",
    "Egypt": "EGY", "Gambia": "GMB", "Gaza Strip and West Bank": "PSE",
    "Hong Kong SAR": "HKG", "Iran": "IRN", "Kyrgyzstan": "KGZ", "Laos": "LAO",
    "Korea, Democratic People's Republic of (North)": "PRK",
    "Korea, Republic of (South)": "KOR", "Macau": "MAC",
    "Micronesia, Federated States of": "FSM", "Slovakia": "SVK", "Somalia": "SOM",
    "St Kitts and Nevis": "KNA", "St Lucia": "LCA", "St Vincent and the Grenadines": "VCT",
    "Swaziland": "SWZ", "Syria": "SYR", "Taiwan": "TWN", "United States of America": "USA",
    "Venezuela": "VEN", "Vietnam": "VNM", "Virgin Islands, British": "VGB", "Yemen": "YEM",
    "Samoa, American": "ASM", "Puerto Rico": "PRI", "Nauru": "NRU", "Anguilla": "AIA",
    "Cook Islands": "COK", "Falkland Islands": "FLK", "Pitcairn Islands": "PCN",
    "St Helena": "SHN", "Reunion": "REU", "Guadeloupe": "GLP", "French Guiana": "GUF",
    "Kosovo": "XKX", "Turkey": "TUR", "Cape Verde": "CPV", "Brunei Darussalam": "BRN",
    "Russian Federation": "RUS", "Macedonia": "MKD", "Burma (Myanmar)": "MMR",
    "Ivory Coast": "CIV", "Cote d'Ivoire": "CIV",
    # Not countries, or countries that no longer exist: excluded from maps.
    "Other": None, "Yugoslavia": None, "Netherlands Antilles": None,
    "United Kingdom (Other)": None, "United States (Territories)": None,
    "Bouvet Island": None, "British Indian Ocean Territory": None,
    "Cocos (Keeling) Islands": None,
}

# The World Bank doesn't publish Taiwan. It sends real numbers of students,
# so its population is entered by hand and cited on the page.
POPULATION_EXTRA = {
    "TWN": (23_400_000, "Taiwan Ministry of the Interior, Dept. of Household Registration, end-2024"),
}

# The map file's "inside the shape" point lands on the wrong island for a few
# countries. These put the circle where most people (and students) are.
POINT_OVERRIDE = {
    "MYS": (101.7, 3.6),     # Peninsular Malaysia, not Borneo
    "IDN": (110.4, -7.3),    # Java, not Sumatra
    "NZL": (175.4, -39.0),   # North Island, not South Island
}

# Countries this small get a per-capita rate, but it's flagged and left out of
# the ranking: 130 students from Nauru (population 12,000) is a rate of 1,100
# per 100k, which is a statistical quirk rather than a story. The cutoff is
# 200,000 rather than 1 million so that Bhutan (800,000 people, 15,000
# enrolments) is ranked: that one is real.
SMALL_POPULATION = 200_000

FLOW_TOP_ORIGINS = 12
FLOW_MIN = 500   # drop flows under 500 enrolments, or the map becomes a hairball


def save(df: pd.DataFrame, name: str) -> None:
    path = DATA / name
    df.to_csv(path, index=False)
    print(f"  wrote {name:<28} {len(df):>5,} rows  {path.stat().st_size / 1024:>6.1f} KB")


def need(path: Path) -> Path:
    if not path.exists():
        raise SystemExit(f"Missing {path.relative_to(ROOT)}. Run the earlier prep scripts / check_setup.py")
    return path


# ------------------------------------------------------------------ loaders
def load_world_bank() -> pd.DataFrame:
    pop = pd.read_csv(need(RAW / "worldbank_population.csv"), skiprows=4)
    meta = pd.read_csv(need(RAW / "worldbank_country_metadata.csv"))
    # Rows with a blank Region are aggregates ("World", "High income", ...).
    countries = set(meta.loc[meta["Region"].notna(), "Country Code"])
    pop = pop[pop["Country Code"].isin(countries)]
    out = pop[["Country Name", "Country Code", POP_YEAR]].rename(
        columns={"Country Name": "wb_name", "Country Code": "iso3", POP_YEAR: "population"})
    out["pop_source"] = f"World Bank SP.POP.TOTL, {POP_YEAR}"
    extra = pd.DataFrame([{"wb_name": None, "iso3": k, "population": v, "pop_source": s}
                          for k, (v, s) in POPULATION_EXTRA.items()])
    return pd.concat([out, extra], ignore_index=True)


def load_abs_states() -> pd.DataFrame:
    """ABS Table 3: estimated resident population by state, Dec 2025."""
    ws = load_workbook(need(RAW / "31010do001_202512.xlsx"), read_only=True)["Table_3"]
    rows = list(ws.iter_rows(values_only=True))
    # Find the year header row, then the column under POPULATION for 2025.
    hdr_i = next(i for i, r in enumerate(rows) if r and int(YEAR) in r)
    col = rows[hdr_i].index(int(YEAR))   # first occurrence = POPULATION block
    names = {"New South Wales", "Victoria", "Queensland", "South Australia", "Western Australia",
             "Tasmania", "Northern Territory", "Australian Capital Territory"}
    recs = [{"name": r[0], "population": int(r[col])} for r in rows[hdr_i + 1:]
            if r and r[0] in names]
    return pd.DataFrame(recs)


# --------------------------------------------------------------------- main
def main() -> None:
    nat_year = pd.read_csv(need(DATA / "nationality_by_year.csv"))
    nat_region = pd.read_csv(need(DATA / "nationality_region.csv"))
    points = pd.read_csv(need(GEO / "country_points.csv"))
    wb = load_world_bank()

    # ---- 1. resolve every Department of Education name to ISO3 -------------
    names = sorted(set(nat_year["nationality"]) | set(nat_region["nationality"]))
    by_wb = dict(zip(wb["wb_name"].dropna(), wb.loc[wb["wb_name"].notna(), "iso3"]))
    by_map = dict(zip(points["name"], points["iso3"]))

    rows = []
    for n in names:
        if n in DOE_TO_ISO3:
            iso, how = DOE_TO_ISO3[n], ("manual" if DOE_TO_ISO3[n] else "excluded")
        elif n in by_wb:
            iso, how = by_wb[n], "World Bank name"
        elif n in by_map:
            iso, how = by_map[n], "map name"
        else:
            iso, how = None, "UNMATCHED"
        rows.append({"nationality": n, "iso3": iso, "matched_by": how})
    lookup = pd.DataFrame(rows)

    lookup = lookup.merge(nat_region[["nationality", "region_group"]], on="nationality", how="left")
    pts = points.set_index("iso3")[["lon", "lat"]].copy()
    for iso, (lon, lat) in POINT_OVERRIDE.items():
        pts.loc[iso] = [lon, lat]
    lookup = lookup.merge(pts, left_on="iso3", right_index=True, how="left")
    lookup = lookup.merge(wb[["iso3", "population", "pop_source"]], on="iso3", how="left")

    # A few old names (e.g. Zaire) have no region in the 2019-25 workbook;
    # borrow the region of the other name that maps to the same ISO3.
    fill = lookup.dropna(subset=["region_group"]).drop_duplicates("iso3").set_index("iso3")["region_group"]
    lookup["region_group"] = lookup["region_group"].fillna(lookup["iso3"].map(fill))

    print("\nWriting CSVs")
    save(lookup, "country_lookup.csv")

    # ---- 2. Maps 1 and 2: one row per country, latest year -----------------
    latest = nat_year[nat_year["year"] == YEAR].merge(lookup, on="nationality", how="left")
    country = (latest.dropna(subset=["iso3"])
                     .groupby("iso3", as_index=False)
                     .agg(country=("nationality", "first"), enrolments=("enrolments", "sum"),
                          region_group=("region_group", "first"), lat=("lat", "first"),
                          lon=("lon", "first"), population=("population", "first")))
    country["per_100k"] = (country["enrolments"] / country["population"] * 100_000).round(2)
    country["small_population"] = country["population"] < SMALL_POPULATION
    country["rank_enrolments"] = country["enrolments"].rank(ascending=False, method="min").astype(int)
    big = country[~country["small_population"] & country["per_100k"].notna()]
    country["rank_per_100k"] = big["per_100k"].rank(ascending=False, method="min")
    country["rank_per_100k"] = country["rank_per_100k"].astype("Int64")
    country = country.sort_values("enrolments", ascending=False)
    save(country, "country_2025.csv")

    # ---- 3. Map 4 fill: states ----------------------------------------------
    caps = pd.read_csv(need(GEO / "au_state_capitals.csv"))
    st_enrol = pd.read_csv(need(DATA / "state_by_year.csv"))
    st_enrol = st_enrol[st_enrol["year"] == YEAR][["state", "enrolments"]]
    states = caps.merge(load_abs_states(), on="name").merge(st_enrol, on="state")
    states["per_1000"] = (states["enrolments"] / states["population"] * 1000).round(1)
    states["share"] = (states["enrolments"] / states["enrolments"].sum()).round(4)
    save(states.sort_values("enrolments", ascending=False), "state_2025.csv")

    # ---- 4. Map 3: flows, origin -> state capital ---------------------------
    flows = pd.read_csv(need(DATA / "nationality_state_2025.csv"))
    top = country.head(FLOW_TOP_ORIGINS)["country"].tolist()
    flows = flows[flows["nationality"].isin(top) & (flows["enrolments"] >= FLOW_MIN)]
    flows = (flows.merge(lookup[["nationality", "iso3", "region_group", "lat", "lon"]], on="nationality")
                  .merge(caps[["state", "city", "lat", "lon"]], on="state", suffixes=("", "_dest"))
                  .rename(columns={"lat": "origin_lat", "lon": "origin_lon",
                                   "lat_dest": "dest_lat", "lon_dest": "dest_lon"}))
    # Share of that origin's enrolments going to this state: the number the
    # flow map's argument rests on ("most Indian students go to Victoria").
    origin_tot = country.set_index("country")["enrolments"]
    flows["share_of_origin"] = (flows["enrolments"] / flows["nationality"].map(origin_tot)).round(4)
    save(flows.sort_values("enrolments", ascending=False), "flows_2025.csv")

    # ---- 5. Map 4 circles + beeswarm: providers with coordinates -----------
    inst = pd.read_csv(need(DATA / "he_institution_2024.csv"))
    locs = pd.read_csv(need(GEO / "au_university_locations.csv"))
    inst_geo = inst.merge(locs, on="institution", how="left")
    save(inst_geo, "he_institution_geo_2024.csv")

    validate(lookup, latest, country, states, flows, inst_geo)


# --------------------------------------------------------------- validation
def validate(lookup, latest, country, states, flows, inst_geo) -> None:
    print("\nValidation")
    failures = 0

    def check(label, ok, detail=""):
        nonlocal failures
        failures += not ok
        print(f"  [{' OK ' if ok else 'FAIL'}] {label}{(': ' + detail) if detail else ''}")

    unmatched = lookup.loc[lookup["matched_by"] == "UNMATCHED", "nationality"].tolist()
    check("every nationality name resolved", not unmatched,
          ("add to DOE_TO_ISO3: " + "; ".join(unmatched)) if unmatched else f"{len(lookup)} names")

    total = latest["enrolments"].sum()
    mapped = latest.dropna(subset=["iso3"])["enrolments"].sum()
    check(f"{YEAR} enrolments placed on a country", mapped / total > 0.999,
          f"{mapped / total:.3%} ({total - mapped:,.0f} in 'Other' and similar)")

    no_pt = country[country["lat"].isna()]
    check("every country with students has a map point", no_pt.empty,
          ", ".join(f"{r.country} ({r.enrolments:,})" for r in no_pt.itertuples()) or f"{len(country)} countries")
    no_pop = country[country["population"].isna()]
    print(f"  [info] no population figure (left grey on Map 2): "
          + (", ".join(f"{r.country} ({r.enrolments:,})" for r in no_pop.itertuples()) or "none"))

    top = country.dropna(subset=["rank_per_100k"]).sort_values("rank_per_100k").head(5)
    print(f"  [info] highest rates, countries over {SMALL_POPULATION:,} people:")
    for r in top.itertuples():
        print(f"           {int(r.rank_per_100k)}. {r.country:<12} {r.per_100k:>7.1f} per 100k  ({r.enrolments:,} students)")
    small = country[country["small_population"] & (country["per_100k"] > top["per_100k"].iloc[0])]
    if len(small):
        print("  [info] small nations with higher rates (flagged small_population): "
              + ", ".join(f"{r.country} {r.per_100k:.0f}" for r in small.itertuples()))
    for c in ("China", "India"):
        r = country[country["country"] == c].iloc[0]
        print(f"  [info] {c}: {r.per_100k:.1f} per 100k (rank {r.rank_per_100k})")

    check("8 states with ABS population", len(states) == 8 and states["population"].notna().all())
    for r in states.itertuples():
        print(f"           {r.state:<4} {r.enrolments:>8,} enrolments  {r.per_1000:>5.1f} per 1,000 residents")

    check("flows reach several different states", flows["state"].nunique() >= 5,
          f"{len(flows)} flows into {flows['state'].nunique()} states")
    missing = inst_geo.loc[inst_geo["lat"].isna(), "institution"].tolist()
    check("every provider has campus coordinates", not missing, "; ".join(missing) or f"{len(inst_geo)}")

    print()
    print("All checks passed." if failures == 0 else f"{failures} check(s) failed.")
    sys.exit(1 if failures else 0)


if __name__ == "__main__":
    main()
